"""
Analah Capital – Daily Ops & Task Infrastructure
- Separate downloadable POA and EOD reports
- One Master Excel file for current tasks (team-wise)
"""

import streamlit as st
import pandas as pd
from datetime import date
import io
import os
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from utils.auth import (
    init_auth_session,
    check_login,
    is_admin,
    logout,
    get_current_user,
)
from utils.data import (
    add_entry,
    update_status,
    delete_entry,
    get_user_entries,
    get_entries_by_date,
    get_open_tasks,
    get_all_data,
    get_master_file_path,
)

# -------------------- PAGE CONFIG --------------------
st.set_page_config(
    page_title="Analah Capital | Daily Ops",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# -------------------- AUTH --------------------
init_auth_session()

if not st.session_state.logged_in:
    st.markdown(
        """
        <div style='text-align:center; padding: 2.5rem 0 1rem 0;'>
            <h1 style='color:#1E88E5; margin-bottom:0.2rem;'>Analah Capital</h1>
            <h3 style='color:#555; font-weight:400;'>Daily Operations & Task Infrastructure</h3>
        </div>
        """,
        unsafe_allow_html=True,
    )

    col1, col2, col3 = st.columns([1, 1.3, 1])
    with col2:
        with st.form("login_form"):
            username = st.text_input("Username", placeholder="Enter username")
            password = st.text_input("Password", type="password", placeholder="Enter password")
            submitted = st.form_submit_button("Login", use_container_width=True, type="primary")

            if submitted:
                user = check_login(username, password)
                if user:
                    st.session_state.logged_in = True
                    st.session_state.username = user["username"]
                    st.session_state.user_info = user
                    st.rerun()
                else:
                    st.error("Invalid username or password")
    st.stop()

# -------------------- LOGGED-IN STATE --------------------
user = get_current_user()
username = st.session_state.username
admin = is_admin(user)
today = date.today().isoformat()

# -------------------- SIDEBAR --------------------
with st.sidebar:
    st.markdown(f"### 👋 {user['name']}")
    st.caption(f"Role: {'Administrator' if admin else 'Team Member'}")
    st.divider()

    if admin:
        nav_options = [
            "📝 My POA / EOD",
            "📋 My Tasks",
            "📊 Admin Dashboard",
            "📥 Downloads & Reports",
            "📜 Activity Log",
        ]
    else:
        nav_options = [
            "📝 My POA / EOD",
            "📋 My Tasks",
            "📜 My History",
        ]

    page = st.radio("Navigation", nav_options, label_visibility="collapsed")
    st.divider()

    if st.button("Logout", use_container_width=True):
        logout()


# -------------------- HELPER: Create Excel bytes --------------------
def create_excel_bytes(df: pd.DataFrame) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Report")
        ws = writer.sheets["Report"]

        header_fill = PatternFill(start_color="1E88E5", end_color="1E88E5", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        thin = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = thin
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    max_len = max(max_len, len(str(cell.value or "")))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 3, 55)

    output.seek(0)
    return output.getvalue()


# -------------------- PAGES --------------------

# ========== 1. MY POA / EOD ==========
if page == "📝 My POA / EOD":
    st.title("Daily Plan of Action & EOD")
    st.caption(f"Date: **{today}**")

    tab_poa, tab_eod = st.tabs(["Plan of Action (POA)", "End of Day (EOD)"])

    with tab_poa:
        with st.form("poa_form", clear_on_submit=True):
            content = st.text_area(
                "Today's Plan of Action",
                height=180,
                placeholder="Key priorities, meetings, deals, follow-ups, client calls...",
            )
            priority = st.selectbox("Priority", ["High", "Medium", "Low"])
            if st.form_submit_button("Submit POA", type="primary"):
                if content.strip():
                    add_entry(
                        username=username,
                        name=user["name"],
                        entry_type="POA",
                        content=content,
                        priority=priority,
                        status="Open",
                    )
                    st.success("POA submitted successfully!")
                    st.rerun()
                else:
                    st.warning("Please enter some content.")

    with tab_eod:
        with st.form("eod_form", clear_on_submit=True):
            content = st.text_area(
                "End of Day Summary / Completed Work",
                height=180,
                placeholder="What was completed? Key outcomes, blockers, handovers, next steps...",
            )
            if st.form_submit_button("Submit EOD", type="primary"):
                if content.strip():
                    add_entry(
                        username=username,
                        name=user["name"],
                        entry_type="EOD",
                        content=content,
                        priority="Medium",
                        status="Completed",
                    )
                    st.success("EOD submitted!")
                    st.rerun()
                else:
                    st.warning("Please enter some content.")

# ========== 2. MY TASKS ==========
elif page == "📋 My Tasks":
    st.title("My Current Tasks")

    with st.expander("➕ Add New Task", expanded=False):
        with st.form("add_task_form"):
            t_content = st.text_input("Task description")
            t_priority = st.selectbox("Priority", ["High", "Medium", "Low"])
            if st.form_submit_button("Add Task"):
                if t_content.strip():
                    add_entry(
                        username=username,
                        name=user["name"],
                        entry_type="Task",
                        content=t_content,
                        priority=t_priority,
                        status="Open",
                    )
                    st.rerun()
                else:
                    st.warning("Task description cannot be empty.")

    my_df = get_user_entries(username)

    if my_df.empty:
        st.info("No tasks yet. Add your first one above.")
    else:
        for _, row in my_df.iterrows():
            with st.container(border=True):
                c1, c2, c3 = st.columns([0.68, 0.18, 0.14])
                c1.markdown(f"**{row['content']}**")
                c1.caption(
                    f"{row['type']} • {row['priority']} • {row['date']} • "
                    f"Last update: {str(row['updated_at'])[:16]}"
                )

                status_options = ["Open", "In Progress", "Completed", "Blocked"]
                current_idx = (
                    status_options.index(row["status"])
                    if row["status"] in status_options
                    else 0
                )
                new_status = c2.selectbox(
                    "Status",
                    status_options,
                    index=current_idx,
                    key=f"status_{row['id']}",
                    label_visibility="collapsed",
                )
                if new_status != row["status"]:
                    update_status(row["id"], new_status)
                    st.rerun()

                if c3.button("🗑️", key=f"del_{row['id']}", help="Delete task"):
                    delete_entry(row["id"])
                    st.rerun()

# ========== 3. ADMIN DASHBOARD ==========
elif page == "📊 Admin Dashboard" and admin:
    st.title("Admin Dashboard – Current Tasks")
    st.caption("Live view of all open tasks from the Master Excel file (team-wise)")

    open_df = get_open_tasks()

    if open_df.empty:
        st.success("No open tasks right now 🎉")
    else:
        st.dataframe(
            open_df[["date", "name", "type", "content", "priority", "status", "updated_at"]],
            use_container_width=True,
            hide_index=True,
        )

        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Open", len(open_df[open_df["status"] == "Open"]))
        m2.metric("In Progress", len(open_df[open_df["status"] == "In Progress"]))
        m3.metric("High Priority", len(open_df[open_df["priority"] == "High"]))
        m4.metric("Active Members", open_df["name"].nunique())

# ========== 4. DOWNLOADS & REPORTS ==========
elif page == "📥 Downloads & Reports" and admin:
    st.title("Downloads & Reports")

    # ---------- Daily POA / EOD Reports ----------
    st.subheader("1. Daily POA / EOD Reports")
    st.caption("Download separate Excel files for Plan of Action or End of Day")

    c1, c2 = st.columns(2)
    with c1:
        report_date = st.date_input("Select Date", value=date.today())
    with c2:
        report_type = st.selectbox("Report Type", ["POA", "EOD"])

    filtered = get_entries_by_date(report_date.isoformat(), entry_type=report_type)

    if filtered.empty:
        st.warning(f"No {report_type} records found for {report_date.isoformat()}")
    else:
        st.dataframe(filtered, use_container_width=True, hide_index=True)

        filename = f"{report_date.isoformat()}_{report_type}.xlsx"
        excel_bytes = create_excel_bytes(filtered)

        st.download_button(
            label=f"⬇️ Download {filename}",
            data=excel_bytes,
            file_name=filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )

    st.divider()

    # ---------- Master Tasks File ----------
    st.subheader("2. Master Tasks File (Team-wise)")
    st.caption("This is the single existing file that contains all current tasks of the team.")

    master_path = get_master_file_path()
    if os.path.exists(master_path):
        with open(master_path, "rb") as f:
            st.download_button(
                label="⬇️ Download Master Tasks Excel (All Current Tasks)",
                data=f,
                file_name="Analah_Master_Tasks.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
            )
    else:
        st.info("Master file will be created automatically after the first entry is submitted.")

# ========== 5. ACTIVITY / HISTORY ==========
elif page in ["📜 Activity Log", "📜 My History"]:
    st.title("Activity Log" if admin else "My History")

    view_df = get_all_data() if admin else get_user_entries(username)

    if view_df.empty:
        st.info("No activity recorded yet.")
    else:
        st.dataframe(view_df, use_container_width=True, hide_index=True)

# -------------------- FOOTER --------------------
st.markdown("---")
st.caption("Analah Capital Internal Ops Platform • Light Theme • Master Excel Storage")