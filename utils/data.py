"""
Data layer for Analah Capital Ops Platform
- Master Excel file for all tasks (team-wise)
- Supports Team column
"""

import os
import pandas as pd
from datetime import datetime, date
from typing import Optional
import streamlit as st
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DATA_DIR = "data"
MASTER_FILE = os.path.join(DATA_DIR, "Analah_Master_Tasks.xlsx")

COLUMNS = [
    "id",
    "date",
    "username",
    "name",
    "team",           # NEW
    "type",           # POA | EOD | Task
    "content",
    "priority",
    "status",
    "created_at",
    "updated_at",
]


def ensure_data_dir():
    os.makedirs(DATA_DIR, exist_ok=True)


def _style_worksheet(ws):
    header_fill = PatternFill(start_color="1E88E5", end_color="1E88E5", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 3, 55)


def load_data() -> pd.DataFrame:
    ensure_data_dir()
    if os.path.exists(MASTER_FILE):
        try:
            df = pd.read_excel(MASTER_FILE, engine="openpyxl")
            for col in COLUMNS:
                if col not in df.columns:
                    df[col] = None
            return df[COLUMNS]
        except Exception as e:
            st.warning(f"Could not read master Excel: {e}. Starting fresh.")
    return pd.DataFrame(columns=COLUMNS)


def save_data(df: pd.DataFrame) -> None:
    ensure_data_dir()
    df = df[COLUMNS].dropna(how="all")
    with pd.ExcelWriter(MASTER_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="All_Data")
        ws = writer.sheets["All_Data"]
        _style_worksheet(ws)


def generate_id(username: str) -> str:
    return f"{username}_{datetime.now().timestamp():.6f}"


def add_entry(
    username: str,
    name: str,
    team: str,
    entry_type: str,
    content: str,
    priority: str = "Medium",
    status: str = "Open",
    entry_date: Optional[str] = None,
) -> pd.DataFrame:
    df = load_data()
    now = datetime.now().isoformat()

    new_row = {
        "id": generate_id(username),
        "date": entry_date or date.today().isoformat(),
        "username": username,
        "name": name,
        "team": team,
        "type": entry_type,
        "content": content.strip(),
        "priority": priority,
        "status": status,
        "created_at": now,
        "updated_at": now,
    }

    df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
    save_data(df)
    return df


def update_status(entry_id: str, new_status: str) -> pd.DataFrame:
    df = load_data()
    mask = df["id"] == entry_id
    if mask.any():
        df.loc[mask, "status"] = new_status
        df.loc[mask, "updated_at"] = datetime.now().isoformat()
        save_data(df)
    return df


def delete_entry(entry_id: str) -> pd.DataFrame:
    df = load_data()
    df = df[df["id"] != entry_id].reset_index(drop=True)
    save_data(df)
    return df


def get_user_entries(username: str, entry_type: Optional[str] = None) -> pd.DataFrame:
    df = load_data()
    mask = df["username"] == username
    if entry_type:
        mask = mask & (df["type"] == entry_type)
    return df[mask].sort_values("created_at", ascending=False)


def get_entries_by_date(
    target_date: str,
    entry_type: Optional[str] = None,
    team: Optional[str] = None,
) -> pd.DataFrame:
    df = load_data()
    mask = df["date"] == target_date
    if entry_type:
        mask = mask & (df["type"] == entry_type)
    if team and team != "All Teams":
        mask = mask & (df["team"] == team)
    return df[mask].sort_values(["team", "name", "created_at"])


def get_open_tasks(team: Optional[str] = None) -> pd.DataFrame:
    df = load_data()
    mask = df["status"].isin(["Open", "In Progress"])
    if team and team != "All Teams":
        mask = mask & (df["team"] == team)
    return df[mask].sort_values(["team", "date", "priority"], ascending=[True, False, True])


def get_all_data() -> pd.DataFrame:
    return load_data().sort_values("created_at", ascending=False)


def get_master_file_path() -> str:
    return MASTER_FILE
